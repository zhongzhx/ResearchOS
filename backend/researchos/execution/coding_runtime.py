from __future__ import annotations

import csv
import html
import importlib.util
import json
import secrets
import statistics
import struct
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path
from typing import Any

from backend.researchos.execution.runtime_adapter import redact_text
from backend.researchos.workspace.file_artifact_registry import FileArtifactRegistry
from backend.researchos.workspace.project_workspace import ProjectWorkspace


def _is_under(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _safe_segment(value: str, fallback: str) -> str:
    normalized = "".join(char if char.isalnum() or char in "._-" else "-" for char in str(value or "")).strip(" .-_")
    return normalized or fallback


def _preview_type(path: Path) -> str:
    return {
        ".md": "markdown",
        ".csv": "table",
        ".svg": "figure",
        ".png": "figure",
        ".pptx": "pptx",
        ".json": "data",
    }.get(path.suffix.lower(), "file")


class CodingRuntimeService:
    def __init__(self, agent_root: Path) -> None:
        self.agent_root = Path(agent_root).resolve()
        self.registry = FileArtifactRegistry(self.agent_root)

    def create_run_workspace(self, project_id: str, run_id: str = "") -> dict[str, Path]:
        self.registry._project(project_id)
        safe_run_id = _safe_segment(run_id, f"coding_{secrets.token_hex(8)}")
        root = ProjectWorkspace(self.agent_root, project_id).root / "runs" / safe_run_id
        paths = {"run_root": root, "workspace": root / "workspace", "outputs": root / "outputs", "logs": root / "logs"}
        for path in paths.values():
            path.mkdir(parents=True, exist_ok=True)
        return paths

    def resolve_output_path(self, paths: dict[str, Path], filename: str) -> Path:
        raw = str(filename or "")
        if not raw or Path(raw).name != raw:
            raise ValueError("output path escapes run outputs directory")
        output = (paths["outputs"] / raw).resolve()
        if not _is_under(output, paths["outputs"]):
            raise ValueError("output path escapes run outputs directory")
        return output

    def _require_dependency(self, module_name: str, package_name: str = "") -> None:
        if importlib.util.find_spec(module_name) is None:
            raise ModuleNotFoundError(package_name or module_name)

    def _register_outputs(self, project_id: str, run_id: str, outputs: list[tuple[Path, str]]) -> list[dict[str, Any]]:
        artifacts = []
        for path, source_type in outputs:
            artifact = self.registry.register_file(
                project_id=project_id,
                source_type=source_type,
                file_path=path,
                display_name=path.name,
                run_id=run_id,
                status="generated",
                ingest_status="generated",
                metadata={"coding_runtime_run_id": run_id, "preview_type": _preview_type(path)},
            )
            artifacts.append(
                {
                    "artifact_id": artifact["artifact_id"],
                    "display_name": artifact["display_name"],
                    "type": artifact["source_type"],
                    "path": artifact["absolute_path"],
                    "mime_type": artifact["mime_type"],
                    "preview_type": artifact["preview_type"],
                }
            )
        return artifacts

    def _result(
        self,
        *,
        ok: bool,
        status: str,
        run_id: str,
        project_id: str,
        artifacts: list[dict[str, Any]] | None = None,
        stdout_summary: str = "",
        stderr_summary: str = "",
        qa_report: str = "",
        user_message: str = "",
        diagnostics: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        return {
            "ok": ok,
            "status": status,
            "run_id": run_id,
            "project_id": project_id,
            "artifacts": artifacts or [],
            "stdout_summary": redact_text(stdout_summary, max_chars=1200),
            "stderr_summary": redact_text(stderr_summary, max_chars=1200),
            "qa_report": qa_report,
            "user_message": user_message,
            "developer_diagnostics": diagnostics or {},
        }

    def execute_template(self, project_id: str, template: str, params: dict[str, Any], *, run_id: str = "") -> dict[str, Any]:
        paths = self.create_run_workspace(project_id, run_id)
        actual_run_id = paths["run_root"].name
        try:
            handler = getattr(self, f"_template_{template}", None)
            if not handler:
                return self._result(ok=False, status="failed", run_id=actual_run_id, project_id=project_id, stderr_summary=f"unknown template: {template}", user_message="Unsupported coding runtime template.")
            outputs, qa_report = handler(paths, params or {})
            artifacts = self._register_outputs(project_id, actual_run_id, outputs)
            return self._result(
                ok=True,
                status="completed",
                run_id=actual_run_id,
                project_id=project_id,
                artifacts=artifacts,
                stdout_summary=f"Generated {len(artifacts)} artifact(s) with template {template}.",
                qa_report=qa_report,
                user_message=f"Generated {len(artifacts)} project artifact(s).",
                diagnostics={"template": template, "run_root": str(paths["run_root"])},
            )
        except ModuleNotFoundError as exc:
            return self._result(ok=False, status="dependency_missing", run_id=actual_run_id, project_id=project_id, stderr_summary=str(exc), user_message=f"Missing local dependency: {exc.name or exc}.")
        except ValueError as exc:
            return self._result(ok=False, status="needs_input", run_id=actual_run_id, project_id=project_id, stderr_summary=str(exc), user_message=str(exc))
        except Exception as exc:  # noqa: BLE001
            return self._result(ok=False, status="failed", run_id=actual_run_id, project_id=project_id, stderr_summary=str(exc), user_message="Local coding runtime failed.")

    def execute_python_argv(
        self,
        project_id: str,
        argv: list[str],
        *,
        run_id: str = "",
        timeout_seconds: int = 30,
    ) -> dict[str, Any]:
        paths = self.create_run_workspace(project_id, run_id)
        actual_run_id = paths["run_root"].name
        args = [str(item) for item in argv or []]
        try:
            if not args:
                raise ValueError("python argv is required")
            python_names = {"python", "python.exe", "python3", "python3.exe", Path(sys.executable).name.casefold()}
            if Path(args[0]).name.casefold() not in python_names:
                raise ValueError("coding runtime only allows Python argv execution")
            if len(args) < 2 or args[1] in {"-c", "-m"}:
                raise ValueError("coding runtime requires a workspace Python script")
            script = (paths["workspace"] / args[1]).resolve() if not Path(args[1]).is_absolute() else Path(args[1]).resolve()
            if script.suffix.lower() != ".py" or not _is_under(script, paths["workspace"]) or not script.is_file():
                raise ValueError("Python script must exist inside the run workspace")
            for arg in args[2:]:
                if arg.startswith("-"):
                    continue
                value = Path(arg)
                if value.suffix:
                    resolved = (paths["workspace"] / value).resolve() if not value.is_absolute() else value.resolve()
                    if not _is_under(resolved, paths["run_root"]):
                        raise ValueError("Python argv path escapes the run workspace")
            completed = subprocess.run(
                args,
                cwd=str(paths["workspace"]),
                capture_output=True,
                text=True,
                timeout=max(1, min(int(timeout_seconds or 30), 120)),
                check=False,
            )
            (paths["logs"] / "stdout.log").write_text(redact_text(completed.stdout, max_chars=10000), encoding="utf-8")
            (paths["logs"] / "stderr.log").write_text(redact_text(completed.stderr, max_chars=10000), encoding="utf-8")
            if completed.returncode:
                return self._result(
                    ok=False,
                    status="failed",
                    run_id=actual_run_id,
                    project_id=project_id,
                    stdout_summary=completed.stdout,
                    stderr_summary=completed.stderr,
                    user_message="Local Python execution failed.",
                    diagnostics={"exit_code": completed.returncode},
                )
            outputs = [(path, self._source_type_for_output(path)) for path in paths["outputs"].iterdir() if path.is_file()]
            artifacts = self._register_outputs(project_id, actual_run_id, outputs)
            return self._result(
                ok=True,
                status="completed",
                run_id=actual_run_id,
                project_id=project_id,
                artifacts=artifacts,
                stdout_summary=completed.stdout,
                stderr_summary=completed.stderr,
                user_message=f"Generated {len(artifacts)} project artifact(s).",
                diagnostics={"exit_code": completed.returncode},
            )
        except subprocess.TimeoutExpired:
            return self._result(ok=False, status="failed", run_id=actual_run_id, project_id=project_id, stderr_summary="execution timed out", user_message="Local Python execution timed out.")
        except ValueError as exc:
            return self._result(ok=False, status="failed", run_id=actual_run_id, project_id=project_id, stderr_summary=str(exc), user_message=str(exc))

    def _source_type_for_output(self, path: Path) -> str:
        return {
            ".md": "generated_markdown",
            ".pptx": "generated_pptx",
            ".svg": "generated_figure",
            ".png": "generated_figure",
            ".csv": "generated_table",
        }.get(path.suffix.lower(), "generated_data")

    def _template_markdown_report(self, paths: dict[str, Path], params: dict[str, Any]) -> tuple[list[tuple[Path, str]], str]:
        title = str(params.get("title") or "Research Report").strip()
        sections = params.get("sections") if isinstance(params.get("sections"), dict) else {"Summary": params.get("content") or ""}
        lines = [f"# {title}"]
        for heading, content in sections.items():
            lines.extend(["", f"## {heading}", str(content or "")])
        output = self.resolve_output_path(paths, str(params.get("filename") or "report.md"))
        output.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
        return [(output, "generated_markdown")], "Markdown structure generated locally."

    def _read_table(self, input_path: Path) -> tuple[list[str], list[dict[str, str]]]:
        if input_path.suffix.lower() == ".csv":
            with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                return list(reader.fieldnames or []), list(reader)
        if input_path.suffix.lower() in {".xlsx", ".xls"}:
            self._require_dependency("openpyxl")
            from openpyxl import load_workbook

            sheet = load_workbook(input_path, read_only=True, data_only=True).active
            values = list(sheet.values)
            headers = [str(value or "") for value in (values[0] if values else [])]
            return headers, [{header: str(value if value is not None else "") for header, value in zip(headers, row, strict=False)} for row in values[1:]]
        raise ValueError("data_profile supports csv or xlsx input")

    def _input_path(self, params: dict[str, Any]) -> Path:
        artifact_id = str(params.get("artifact_id") or "").strip()
        project_id = str(params.get("project_id") or "").strip()
        if artifact_id and project_id:
            return Path(self.registry.get(project_id, artifact_id, include_deleted=False)["absolute_path"])
        raw = str(params.get("input_file") or params.get("data_file") or "").strip()
        if not raw:
            raise ValueError("input_file is required")
        resolved = Path(raw).expanduser().resolve()
        declared = {str(Path(item).expanduser().resolve()) for item in params.get("input_files") or []}
        if str(resolved) not in declared:
            raise ValueError("input_file must come from artifact registry or declared TaskSpec.input_files")
        if not resolved.is_file():
            raise ValueError("input_file does not exist")
        return resolved

    def _template_data_profile(self, paths: dict[str, Path], params: dict[str, Any]) -> tuple[list[tuple[Path, str]], str]:
        input_path = self._input_path(params)
        headers, rows = self._read_table(input_path)
        if not headers:
            raise ValueError("input table has no columns")
        columns_path = self.resolve_output_path(paths, "columns.csv")
        missing_path = self.resolve_output_path(paths, "missing_values.csv")
        report_path = self.resolve_output_path(paths, "data_profile.md")
        with columns_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["column", "non_empty_count", "unique_count"])
            for header in headers:
                values = [row.get(header, "") for row in rows if row.get(header, "") != ""]
                writer.writerow([header, len(values), len(set(values))])
        with missing_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["column", "missing_count"])
            for header in headers:
                writer.writerow([header, sum(1 for row in rows if row.get(header, "") == "")])
        report_path.write_text(
            "\n".join(["# Data Profile", "", f"- Rows: {len(rows)}", f"- Columns: {len(headers)}", "", "## Columns", *[f"- {header}" for header in headers]]) + "\n",
            encoding="utf-8",
        )
        return [(report_path, "generated_markdown"), (columns_path, "generated_table"), (missing_path, "generated_table")], "Profile created from registered tabular input."

    def _template_basic_stats_plot(self, paths: dict[str, Path], params: dict[str, Any]) -> tuple[list[tuple[Path, str]], str]:
        input_path = self._input_path(params)
        headers, rows = self._read_table(input_path)
        x_col = str(params.get("x") or params.get("x_column") or "").strip()
        y_col = str(params.get("y") or params.get("y_column") or "").strip()
        if x_col not in headers or y_col not in headers:
            raise ValueError("x and y columns are required")
        grouped: dict[str, list[float]] = {}
        for row in rows:
            try:
                grouped.setdefault(row[x_col], []).append(float(row[y_col]))
            except (TypeError, ValueError):
                continue
        if not grouped:
            raise ValueError("no numeric y values were found")
        labels = list(grouped)
        means = [statistics.fmean(grouped[label]) for label in labels]
        summary_path = self.resolve_output_path(paths, "summary.csv")
        with summary_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow([x_col, "count", "mean"])
            for label, mean in zip(labels, means, strict=True):
                writer.writerow([label, len(grouped[label]), mean])
        svg_path = self.resolve_output_path(paths, "figure.svg")
        png_path = self.resolve_output_path(paths, "figure.png")
        if importlib.util.find_spec("matplotlib") is not None:
            import matplotlib

            matplotlib.use("Agg")
            import matplotlib.pyplot as plt

            fig, ax = plt.subplots(figsize=(6, 4))
            ax.bar(labels, means, color="#2a9d8f")
            ax.set_xlabel(x_col)
            ax.set_ylabel(y_col)
            ax.set_title(str(params.get("figure_goal") or "Scientific summary"))
            fig.tight_layout()
            fig.savefig(svg_path)
            fig.savefig(png_path, dpi=180)
            plt.close(fig)
            renderer = "matplotlib"
        else:
            self._write_svg_bar_chart(svg_path, labels, means, x_col, y_col, str(params.get("figure_goal") or "Scientific summary"))
            self._write_png_bar_chart(png_path, means)
            renderer = "stdlib_fallback"
        code_path = self.resolve_output_path(paths, "figure_code.py")
        code_path.write_text(f"# Generated locally by CodingRuntimeService basic_stats_plot.\n# renderer={renderer}\n", encoding="utf-8")
        spec_path = self.resolve_output_path(paths, "figure_spec.md")
        spec_path.write_text(
            "\n".join(
                [
                    "# Figure Specification",
                    "",
                    f"- Plot type: bar chart",
                    f"- X axis: {x_col}",
                    f"- Y axis: {y_col}",
                    f"- Groups rendered: {len(labels)}",
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        qa_path = self.resolve_output_path(paths, "qa_report.md")
        qa_path.write_text(f"# Figure QA\n\n- Numeric values parsed locally.\n- Vector SVG and PNG exports generated.\n- Renderer: {renderer}.\n", encoding="utf-8")
        return [
            (summary_path, "generated_table"),
            (svg_path, "generated_figure"),
            (png_path, "generated_figure"),
            (code_path, "generated_data"),
            (spec_path, "generated_markdown"),
            (qa_path, "generated_markdown"),
        ], str(qa_path)

    def _template_ppt_from_outline(self, paths: dict[str, Path], params: dict[str, Any]) -> tuple[list[tuple[Path, str]], str]:
        outline = params.get("outline")
        if not outline and any(params.get(key) for key in ["artifact_id", "input_file"]):
            input_path = self._input_path(params)
            if input_path.suffix.lower() == ".json":
                outline = json.loads(input_path.read_text(encoding="utf-8"))
            elif input_path.suffix.lower() in {".md", ".markdown"}:
                outline = [
                    {"title": line.lstrip("# ").strip(), "body": ""}
                    for line in input_path.read_text(encoding="utf-8").splitlines()
                    if line.startswith("#")
                ]
        if not isinstance(outline, list) or not outline:
            raise ValueError("outline is required")
        output = self.resolve_output_path(paths, str(params.get("filename") or "journal_club.pptx"))
        if importlib.util.find_spec("pptx") is not None:
            from pptx import Presentation

            deck = Presentation()
            for index, slide_data in enumerate(outline):
                data = slide_data if isinstance(slide_data, dict) else {"title": str(slide_data), "body": ""}
                layout = deck.slide_layouts[0 if index == 0 else 1]
                slide = deck.slides.add_slide(layout)
                slide.shapes.title.text = str(data.get("title") or f"Slide {index + 1}")
                if index > 0 and len(slide.placeholders) > 1:
                    slide.placeholders[1].text = str(data.get("body") or data.get("content") or "")
            deck.save(output)
            renderer = "python-pptx"
        else:
            self._write_minimal_pptx(output, outline)
            renderer = "stdlib_ooxml_fallback"
        return [(output, "generated_pptx")], f"PPTX package generated with {renderer}."

    def _write_svg_bar_chart(self, path: Path, labels: list[str], means: list[float], x_label: str, y_label: str, title: str) -> None:
        width, height = 720, 480
        left, top, chart_width, chart_height = 72, 60, 600, 320
        maximum = max(means) or 1.0
        slot = chart_width / max(1, len(means))
        bars = []
        for index, (label, mean) in enumerate(zip(labels, means, strict=True)):
            bar_height = chart_height * max(0.0, mean) / maximum
            x = left + index * slot + slot * 0.2
            y = top + chart_height - bar_height
            bars.append(f'<rect x="{x:.2f}" y="{y:.2f}" width="{slot * 0.6:.2f}" height="{bar_height:.2f}" fill="#2a9d8f"/>')
            bars.append(f'<text x="{x + slot * 0.3:.2f}" y="{top + chart_height + 24}" text-anchor="middle">{html.escape(label)}</text>')
        path.write_text(
            "\n".join(
                [
                    f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" role="img">',
                    f"<title>{html.escape(title)}</title>",
                    '<rect width="100%" height="100%" fill="white"/>',
                    f'<line x1="{left}" y1="{top + chart_height}" x2="{left + chart_width}" y2="{top + chart_height}" stroke="#334155"/>',
                    f'<line x1="{left}" y1="{top}" x2="{left}" y2="{top + chart_height}" stroke="#334155"/>',
                    *bars,
                    f'<text x="{left + chart_width / 2}" y="455" text-anchor="middle">{html.escape(x_label)}</text>',
                    f'<text x="18" y="{top + chart_height / 2}" transform="rotate(-90 18,{top + chart_height / 2})" text-anchor="middle">{html.escape(y_label)}</text>',
                    "</svg>",
                ]
            ),
            encoding="utf-8",
        )

    def _write_png_bar_chart(self, path: Path, means: list[float]) -> None:
        width, height = 360, 240
        maximum = max(means) or 1.0
        pixels = bytearray()
        for y in range(height):
            pixels.append(0)
            for x in range(width):
                color = (255, 255, 255)
                for index, mean in enumerate(means):
                    slot = width / max(1, len(means))
                    x0, x1 = int(index * slot + slot * 0.2), int(index * slot + slot * 0.8)
                    y0 = height - int((height - 24) * max(0.0, mean) / maximum)
                    if x0 <= x <= x1 and y0 <= y < height:
                        color = (42, 157, 143)
                        break
                pixels.extend(color)

        def chunk(kind: bytes, data: bytes) -> bytes:
            return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

        path.write_bytes(
            b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(bytes(pixels), 9))
            + chunk(b"IEND", b"")
        )

    def _write_minimal_pptx(self, path: Path, outline: list[Any]) -> None:
        slides = [item if isinstance(item, dict) else {"title": str(item), "body": ""} for item in outline]

        def slide_xml(index: int, item: dict[str, Any]) -> str:
            title = html.escape(str(item.get("title") or f"Slide {index}"))
            body = html.escape(str(item.get("body") or item.get("content") or ""))
            return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree>
    <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>
    <p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="2800" b="1"/><a:t>{title}</a:t></a:r></a:p></p:txBody></p:sp>
    <p:sp><p:nvSpPr><p:cNvPr id="3" name="Body"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1800"/><a:t>{body}</a:t></a:r></a:p></p:txBody></p:sp>
  </p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>"""

        overrides = "".join(f'<Override PartName="/ppt/slides/slide{index}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>' for index in range(1, len(slides) + 1))
        rels = "".join(f'<Relationship Id="rId{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{index}.xml"/>' for index in range(1, len(slides) + 1))
        ids = "".join(f'<p:sldId id="{255 + index}" r:id="rId{index}"/>' for index in range(1, len(slides) + 1))
        files = {
            "[Content_Types].xml": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>{overrides}</Types>""",
            "_rels/.rels": """<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/></Relationships>""",
            "ppt/presentation.xml": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sldIdLst>{ids}</p:sldIdLst><p:sldSz cx="12192000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>""",
            "ppt/_rels/presentation.xml.rels": f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>""",
        }
        for index, item in enumerate(slides, start=1):
            files[f"ppt/slides/slide{index}.xml"] = slide_xml(index, item)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as deck:
            for name, content in files.items():
                deck.writestr(name, content)

    def _template_qpcr_template(self, paths: dict[str, Path], params: dict[str, Any]) -> tuple[list[tuple[Path, str]], str]:
        input_path = self._input_path(params)
        headers, rows = self._read_table(input_path)
        required = {"sample", "target_cq", "reference_cq"}
        if not required.issubset({header.casefold() for header in headers}):
            raise ValueError("qPCR input requires sample, target_cq, and reference_cq columns")
        output = self.resolve_output_path(paths, "qpcr_delta_ct.csv")
        normalized_headers = {header.casefold(): header for header in headers}
        group_header = normalized_headers.get("group")
        calculated: list[dict[str, Any]] = []
        for row in rows:
            try:
                calculated.append(
                    {
                        "sample": row[normalized_headers["sample"]],
                        "group": row.get(group_header, "") if group_header else "",
                        "delta_ct": float(row[normalized_headers["target_cq"]]) - float(row[normalized_headers["reference_cq"]]),
                    }
                )
            except (TypeError, ValueError):
                continue
        if not calculated:
            raise ValueError("qPCR input has no valid numeric Cq rows")
        control_group = str(params.get("control_group") or "").strip()
        if control_group and not group_header:
            raise ValueError("qPCR control_group requires a group column")
        baseline_rows = [item["delta_ct"] for item in calculated if not control_group or item["group"] == control_group]
        if not baseline_rows:
            raise ValueError("qPCR control_group was not found in the input table")
        baseline = statistics.fmean(baseline_rows) if control_group else calculated[0]["delta_ct"]
        for item in calculated:
            item["delta_delta_ct"] = item["delta_ct"] - baseline
            item["relative_expression"] = 2 ** (-item["delta_delta_ct"])
        with output.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["sample", "group", "delta_ct", "delta_delta_ct", "relative_expression"])
            for item in calculated:
                writer.writerow([item["sample"], item["group"], item["delta_ct"], item["delta_delta_ct"], item["relative_expression"]])
        report = self.resolve_output_path(paths, "qpcr_statistics_advice.md")
        baseline_description = f"mean delta Ct for control group '{control_group}'" if control_group else "first valid sample delta Ct; select control_group for confirmatory analysis"
        report.write_text(
            "\n".join(
                [
                    "# qPCR Statistics Advice",
                    "",
                    f"- Delta-delta Ct baseline: {baseline_description}.",
                    "- Relative expression is calculated as 2^(-delta-delta Ct).",
                    "- Confirm biological replicates, technical replicate handling, and the planned statistical test before reporting significance.",
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        labels = [str(item["sample"]) for item in calculated]
        values = [float(item["relative_expression"]) for item in calculated]
        svg_path = self.resolve_output_path(paths, "qpcr_expression.svg")
        png_path = self.resolve_output_path(paths, "qpcr_expression.png")
        self._write_svg_bar_chart(svg_path, labels, values, "sample", "relative expression", "qPCR relative expression draft")
        self._write_png_bar_chart(png_path, values)
        return [
            (output, "generated_table"),
            (report, "generated_markdown"),
            (svg_path, "generated_figure"),
            (png_path, "generated_figure"),
        ], "qPCR delta Ct, delta-delta Ct, relative expression, and draft figure generated."
